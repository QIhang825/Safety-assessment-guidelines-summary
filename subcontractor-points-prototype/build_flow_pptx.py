# -*- coding: utf-8 -*-
"""Build a one-page editable PowerPoint flowchart for subcontractor points."""
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

OUT = r"d:\cursor\cursorproject\subcontractor-points-prototype\分包商安全积分考核-产品流程图.pptx"

NAV = RGBColor(0x00, 0x22, 0x66)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TEXT = RGBColor(0x33, 0x33, 0x33)
MUTED = RGBColor(0x66, 0x66, 0x66)
PERSON_F, PERSON_L = RGBColor(0xE8, 0xF3, 0xFF), RGBColor(0x26, 0x8D, 0xFF)
SUB_F, SUB_L = RGBColor(0xEC, 0xFD, 0xF5), RGBColor(0x11, 0xBC, 0x11)
EXT_F, EXT_L = RGBColor(0xF3, 0xF3, 0xF3), RGBColor(0x8C, 0x8C, 0x8C)
DEC_F, DEC_L = RGBColor(0xFF, 0xF7, 0xE8), RGBColor(0xF7, 0x71, 0x12)
DANGER_F, DANGER_L = RGBColor(0xFF, 0xF1, 0xF0), RGBColor(0xFF, 0x41, 0x41)
PLUS_F = RGBColor(0xF0, 0xF7, 0xFF)
MGR_F, MGR_L = RGBColor(0xF7, 0xF7, 0xFF), RGBColor(0x5B, 0x6C, 0xFF)
LANE_P = RGBColor(0xF7, 0xFB, 0xFF)
LANE_S = RGBColor(0xF6, 0xFD, 0xF8)
LANE_M = RGBColor(0xF7, 0xF7, 0xFF)
LANE_E = RGBColor(0xFA, 0xFA, 0xFA)
BORDER = RGBColor(0xE8, 0xE8, 0xE8)
SLATE = RGBColor(0x47, 0x55, 0x69)


def emu(inches):
    return Inches(inches)


def set_run_font(run, name, size_pt, bold=False, color=TEXT):
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = rPr.find(qn(f"a:{tag}"))
        if el is None:
            el = etree.SubElement(rPr, qn(f"a:{tag}"))
        el.set("typeface", name)


def clear_effects(shape):
    spPr = shape._element.spPr
    el = spPr.find(qn("a:effectLst"))
    if el is not None:
        spPr.remove(el)
    empty = etree.SubElement(spPr, qn("a:effectLst"))
    empty.set("{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst", "")
    # keep empty effectLst to suppress theme shadow
    if empty.getparent() is not None:
        pass


def fill_line(shape, fill, line, weight=1.0):
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(weight)
    clear_effects(shape)


def set_round(shape, adj=0.08):
    try:
        shape.adjustments[0] = adj
    except Exception:
        pass


def add_text(shape, lines, size=10, sub_size=8, bold_first=True, color=TEXT, sub_color=MUTED, align=PP_ALIGN.CENTER):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.03)
    try:
        shape.text_frame.paragraphs[0].alignment = align
    except Exception:
        pass
    tf.anchor = MSO_ANCHOR.MIDDLE
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        p.line_spacing = 1.05
        run = p.add_run()
        run.text = line
        is_title = i == 0
        set_run_font(run, "Microsoft YaHei", size if is_title else sub_size, bold=is_title and bold_first, color=color if is_title else sub_color)


def box(slide, x, y, w, h, lines, fill, line, round_adj=0.08):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, emu(x), emu(y), emu(w), emu(h))
    fill_line(sh, fill, line, 1.15)
    set_round(sh, round_adj)
    add_text(sh, lines, 10, 8)
    return sh


def diamond(slide, x, y, w, h, lines):
    sh = slide.shapes.add_shape(MSO_SHAPE.DIAMOND, emu(x), emu(y), emu(w), emu(h))
    fill_line(sh, DEC_F, DEC_L, 1.25)
    add_text(sh, lines, 9, 8)
    return sh


def lane(slide, x, y, w, h, fill, line, label, label_color):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, emu(x), emu(y), emu(w), emu(h))
    fill_line(sh, fill, line, 0.9)
    set_round(sh, 0.02)
    # lane title as separate text box so boxes stay independent
    tb = slide.shapes.add_textbox(emu(x + 0.08), emu(y + 0.05), emu(1.55), emu(0.28))
    tf = tb.text_frame
    tf.word_wrap = False
    tf.margin_left = Inches(0)
    tf.margin_top = Inches(0)
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = label
    set_run_font(run, "Microsoft YaHei", 10, True, label_color)
    return sh


def arrow(slide, x1, y1, x2, y2, color, dashed=False, weight=1.15):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, emu(x1), emu(y1), emu(x2), emu(y2))
    conn.line.color.rgb = color
    conn.line.width = Pt(weight)
    if dashed:
        conn.line.dash_style = MSO_LINE.DASH
    ln = conn._element.spPr.find(qn("a:ln"))
    if ln is None:
        ln = etree.SubElement(conn._element.spPr, qn("a:ln"))
    for old in list(ln.findall(qn("a:tailEnd"))):
        ln.remove(old)
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")
    clear_effects(conn)
    return conn


def elbow(slide, points, color, dashed=False, weight=1.15):
    """Draw polyline as successive straight connectors; last segment has arrow."""
    segs = []
    for i in range(len(points) - 1):
        x1, y1 = points[i]
        x2, y2 = points[i + 1]
        last = i == len(points) - 2
        conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, emu(x1), emu(y1), emu(x2), emu(y2))
        conn.line.color.rgb = color
        conn.line.width = Pt(weight)
        if dashed:
            conn.line.dash_style = MSO_LINE.DASH
        if last:
            ln = conn._element.spPr.find(qn("a:ln"))
            if ln is None:
                ln = etree.SubElement(conn._element.spPr, qn("a:ln"))
            tail = etree.SubElement(ln, qn("a:tailEnd"))
            tail.set("type", "triangle")
            tail.set("w", "med")
            tail.set("len", "med")
        clear_effects(conn)
        segs.append(conn)
    return segs


def legend_chip(slide, x, y, fill, line, label):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, emu(x), emu(y), emu(0.18), emu(0.18))
    fill_line(sh, fill, line, 0.9)
    set_round(sh, 0.2)
    tb = slide.shapes.add_textbox(emu(x + 0.22), emu(y - 0.02), emu(1.35), emu(0.22))
    tf = tb.text_frame
    tf.margin_left = Inches(0)
    tf.margin_top = Inches(0)
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = label
    set_run_font(run, "Microsoft YaHei", 8, False, MUTED)


def rc(box_xywh):
    x, y, w, h = box_xywh
    return {
        "l": x, "r": x + w, "t": y, "b": y + h,
        "cx": x + w / 2, "cy": y + h / 2,
        "w": w, "h": h,
    }


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slide_layouts[6]  # blank
    s = prs.slides.add_slide(slide)

    bg = s.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(0xF2, 0xF2, 0xF2)

    # title bar
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, emu(0.46))
    fill_line(bar, NAV, NAV, 0)
    bar.line.fill.background()
    tb = s.shapes.add_textbox(emu(0.18), emu(0.08), emu(8.4), emu(0.32))
    tf = tb.text_frame
    tf.margin_left = Inches(0)
    tf.margin_top = Inches(0)
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = "分包商安全积分考核 · 产品流程图"
    set_run_font(run, "Microsoft YaHei", 16, True, WHITE)
    tb2 = s.shapes.add_textbox(emu(8.5), emu(0.1), emu(4.6), emu(0.28))
    tf = tb2.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    run = p.add_run()
    run.text = "劳务分包商 + 作业人员　赋分100　累计至退场不清零"
    set_run_font(run, "Microsoft YaHei", 9, False, RGBColor(0xC8, 0xE2, 0xFA))

    # legend
    chips = [
        (0.18, EXT_F, EXT_L, "劳务教育"),
        (1.55, PERSON_F, PERSON_L, "作业人员积分"),
        (3.15, SUB_F, SUB_L, "分包商积分"),
        (4.65, DEC_F, DEC_L, "分级判断"),
        (5.95, DANGER_F, DANGER_L, "黑名单/清退"),
        (7.35, MGR_F, MGR_L, "项目管理"),
    ]
    for x, f, l, lab in chips:
        legend_chip(s, x, 0.54, f, l, lab)
    tb = s.shapes.add_textbox(emu(8.85), emu(0.52), emu(4.2), emu(0.22))
    tf = tb.text_frame
    tf.margin_left = Inches(0)
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = "虚线箭头 = 人员行为联动扣分包商分"
    set_run_font(run, "Microsoft YaHei", 8, False, SUB_L)

    # lanes
    lane(s, 0.1, 0.78, 13.12, 0.78, LANE_E, BORDER, "劳务教育（外部）", MUTED)
    lane(s, 0.1, 1.62, 13.12, 2.92, LANE_P, RGBColor(0xC8, 0xE2, 0xFA), "作业人员积分", PERSON_L)
    lane(s, 0.1, 4.60, 13.12, 1.18, LANE_S, RGBColor(0xC6, 0xEB, 0xD0), "分包商积分", SUB_L)
    lane(s, 0.1, 5.84, 13.12, 1.50, LANE_M, RGBColor(0xD6, 0xDB, 0xFF), "项目管理", MGR_L)

    # --- 劳务教育 ---
    e1 = (1.85, 0.90, 2.05, 0.56)
    e2 = (8.85, 0.90, 2.20, 0.56)
    box(s, *e1, ["三级教育合格", "人员 / 分包商进场核验"], EXT_F, EXT_L)
    box(s, *e2, ["停工教育考试", "课程 / 学时 / 成绩同步"], EXT_F, EXT_L)

    # --- 人员 row1 ---
    p1 = (1.85, 1.96, 1.95, 0.62)
    p2 = (4.05, 1.96, 2.15, 0.62)
    p3 = (6.40, 1.96, 1.90, 0.62)
    p4 = (8.50, 1.96, 1.85, 0.62)
    pd = (10.70, 1.86, 1.55, 0.82)
    box(s, *p1, ["同步建立人员账户", "赋分 100，写入进场流水"], PERSON_F, PERSON_L)
    box(s, *p2, ["检查隐患映射扣分", "源头扣分后同步；未映射/E类不入"], PERSON_F, PERSON_L)
    box(s, *p3, ["人工违章记分", "A10 / B8 / C5 / D2 / E0"], PERSON_F, PERSON_L)
    box(s, *p4, ["写入人员积分流水", "累计至退场，不按月清零"], PERSON_F, PERSON_L)
    diamond(s, *pd, ["当前分", "分级判断"])

    # --- 人员 row2 分级 ---
    g1 = (1.85, 2.90, 1.85, 0.56)
    g2 = (4.00, 2.90, 2.20, 0.56)
    g3 = (6.50, 2.90, 2.20, 0.56)
    g4 = (9.00, 2.90, 2.05, 0.56)
    box(s, *g1, ["≥90  正常", "继续累计计分"], PERSON_F, PERSON_L)
    box(s, *g2, ["<90  重点管控", "提醒；85–90 自学任务不加分"], DEC_F, DEC_L)
    box(s, *g3, ["<80  停工培训", "下发通知书，7 天内再教育"], DEC_F, DEC_L)
    box(s, *g4, ["<70  清退", "自动列入人员黑名单"], DANGER_F, DANGER_L)

    # --- 人员 row3 ---
    t1 = (6.50, 3.68, 2.20, 0.64)
    t2 = (9.00, 3.68, 2.05, 0.64)
    t3 = (11.20, 3.68, 1.80, 0.64)
    box(s, *t1, ["人工确认培训加分", "首次合格建议 +10，可再调"], PLUS_F, PERSON_L)
    box(s, *t2, ["二次停工 / 培训超时", "不可复岗加分"], DANGER_F, DANGER_L)
    box(s, *t3, ["人员黑名单库", "无审批；本工程公司范围"], DANGER_F, DANGER_L)

    # --- 分包商 ---
    s1 = (1.85, 4.92, 1.95, 0.64)
    s2 = (4.10, 4.92, 2.35, 0.64)
    s3 = (6.70, 4.92, 2.15, 0.64)
    s4 = (9.10, 4.92, 1.95, 0.64)
    s5 = (11.25, 4.92, 1.75, 0.64)
    box(s, *s1, ["同步建立分包商账户", "赋分 100，本项目累计"], SUB_F, SUB_L)
    box(s, *s2, ["人员联动扣分（按人次）", "违规0.1　培训单0.5　黑名单1"], SUB_F, SUB_L)
    box(s, *s3, ["人工评价扣分", "十五条10 / 十六条5 / 牌参照并行"], SUB_F, SUB_L)
    box(s, *s4, ["第十四条直接列入", "不走普通扣分，无审批"], DANGER_F, DANGER_L)
    box(s, *s5, ["分包商状态", "仅 正常 / 黑名单"], SUB_F, SUB_L)

    # --- 管理 ---
    m1 = (1.85, 6.18, 2.15, 0.72)
    m2 = (4.30, 6.18, 2.05, 0.72)
    m3 = (6.65, 6.18, 2.15, 0.72)
    m4 = (9.10, 6.18, 2.15, 0.72)
    box(s, *m1, ["每月 25 日自动快照", "锁定当时累计分与状态，不重置"], MGR_F, MGR_L)
    box(s, *m2, ["月度考核审核", "通报例会并上报工程公司"], MGR_F, MGR_L)
    box(s, *m3, ["退场冻结", "保留历史分与流水，不再考核"], MGR_F, MGR_L)
    box(s, *m4, ["重新进场", "恢复参与，沿用原累计分"], MGR_F, MGR_L)

    # footer
    ft = s.shapes.add_textbox(emu(1.85), emu(7.08), emu(11.2), emu(0.32))
    tf = ft.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0)
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = "人员状态：正常→重点管控→停工培训→清退/黑名单。分包商不因分数自动改状态。自学不加分。不含集团汇总、发卡、平安班组、欠薪门禁、开工条件。"
    set_run_font(run, "Microsoft YaHei", 8, False, MUTED)

    # connectors
    a = {k: rc(v) for k, v in {
        "e1": e1, "e2": e2, "p1": p1, "p2": p2, "p3": p3, "p4": p4, "pd": pd,
        "g1": g1, "g2": g2, "g3": g3, "g4": g4, "t1": t1, "t2": t2, "t3": t3,
        "s1": s1, "s2": s2, "s3": s3, "s4": s4, "s5": s5,
        "m1": m1, "m2": m2, "m3": m3, "m4": m4,
    }.items()}

    # 教育 -> 人员账户
    arrow(s, a["e1"]["cx"], a["e1"]["b"], a["p1"]["cx"], a["p1"]["t"], EXT_L)
    # 教育 -> 分包商账户（左侧通道，不穿过人员框）
    elbow(s, [
        (a["e1"]["l"], a["e1"]["cy"]),
        (1.50, a["e1"]["cy"]),
        (1.50, a["s1"]["t"] - 0.08),
        (a["s1"]["cx"], a["s1"]["t"] - 0.08),
        (a["s1"]["cx"], a["s1"]["t"]),
    ], EXT_L)
    # 停工教育 -> 培训确认（沿教育带下沿左转到人员第三行，避免穿过分级框）
    elbow(s, [
        (a["e2"]["cx"], a["e2"]["b"]),
        (a["e2"]["cx"], 1.58),
        (1.58, 1.58),
        (1.58, a["t1"]["cy"]),
        (a["t1"]["l"], a["t1"]["cy"]),
    ], EXT_L)

    # 人员横向
    arrow(s, a["p1"]["r"], a["p1"]["cy"], a["p2"]["l"], a["p2"]["cy"], PERSON_L)
    arrow(s, a["p2"]["r"], a["p2"]["cy"], a["p3"]["l"], a["p3"]["cy"], PERSON_L)
    arrow(s, a["p3"]["r"], a["p3"]["cy"], a["p4"]["l"], a["p4"]["cy"], PERSON_L)
    arrow(s, a["p4"]["r"], a["p4"]["cy"], a["pd"]["l"], a["pd"]["cy"], PERSON_L)

    # 分级向下
    bus_y = 2.78
    elbow(s, [(a["pd"]["cx"], a["pd"]["b"]), (a["pd"]["cx"], bus_y), (a["g4"]["cx"], bus_y), (a["g4"]["cx"], a["g4"]["t"])], DEC_L)
    elbow(s, [(a["pd"]["cx"], bus_y), (a["g3"]["cx"], bus_y), (a["g3"]["cx"], a["g3"]["t"])], DEC_L)
    elbow(s, [(a["pd"]["cx"], bus_y), (a["g2"]["cx"], bus_y), (a["g2"]["cx"], a["g2"]["t"])], DEC_L)
    elbow(s, [(a["pd"]["cx"], bus_y), (a["g1"]["cx"], bus_y), (a["g1"]["cx"], a["g1"]["t"])], PERSON_L)

    arrow(s, a["g3"]["cx"], a["g3"]["b"], a["t1"]["cx"], a["t1"]["t"], PERSON_L)
    arrow(s, a["g4"]["cx"], a["g4"]["b"], a["t2"]["cx"], a["t2"]["t"], DANGER_L)
    arrow(s, a["t1"]["r"], a["t1"]["cy"], a["t2"]["l"], a["t2"]["cy"], DANGER_L)
    arrow(s, a["t2"]["r"], a["t2"]["cy"], a["t3"]["l"], a["t3"]["cy"], DANGER_L)

    # 分包商横向
    arrow(s, a["s1"]["r"], a["s1"]["cy"], a["s2"]["l"], a["s2"]["cy"], SUB_L)
    arrow(s, a["s2"]["r"], a["s2"]["cy"], a["s3"]["l"], a["s3"]["cy"], SUB_L)
    arrow(s, a["s3"]["r"], a["s3"]["cy"], a["s4"]["l"], a["s4"]["cy"], DANGER_L)
    arrow(s, a["s4"]["r"], a["s4"]["cy"], a["s5"]["l"], a["s5"]["cy"], SUB_L)

    # 管理横向 + 往返
    arrow(s, a["m1"]["r"], a["m1"]["cy"], a["m2"]["l"], a["m2"]["cy"], MGR_L)
    arrow(s, a["m2"]["r"], a["m2"]["cy"], a["m3"]["l"], a["m3"]["cy"], MGR_L)
    arrow(s, a["m3"]["r"], a["m3"]["cy"], a["m4"]["l"], a["m4"]["cy"], MGR_L)
    elbow(s, [
        (a["m4"]["cx"], a["m4"]["b"]),
        (a["m4"]["cx"], 6.98),
        (a["m3"]["cx"], 6.98),
        (a["m3"]["cx"], a["m3"]["b"]),
    ], MGR_L)

    # 联动虚线
    elbow(s, [
        (a["p3"]["cx"], a["p3"]["b"]),
        (a["p3"]["cx"], 4.72),
        (a["s2"]["cx"], 4.72),
        (a["s2"]["cx"], a["s2"]["t"]),
    ], SUB_L, dashed=True)
    elbow(s, [
        (a["g3"]["l"] + 0.18, a["g3"]["b"]),
        (a["g3"]["l"] + 0.18, 4.72),
        (a["s2"]["l"] + 0.35, 4.72),
        (a["s2"]["l"] + 0.35, a["s2"]["t"]),
    ], SUB_L, dashed=True)
    elbow(s, [
        (a["t3"]["cx"], a["t3"]["b"]),
        (a["t3"]["cx"], 4.72),
        (a["s2"]["r"] - 0.35, 4.72),
        (a["s2"]["r"] - 0.35, a["s2"]["t"]),
    ], SUB_L, dashed=True)

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
